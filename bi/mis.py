"""Excel MIS workbook from exported CSV marts."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MIS_SHEETS = (
    "Summary",
    "Aging",
    "GST by rate",
    "Vendors",
    "Exceptions",
)

INR_FORMAT = "#,##,##0.00"
_INT_FORMAT = "#,##,##0"
_HEADER_FONT = Font(bold=True)


def _parse_decimal(value: str) -> Decimal:
    text = (value or "").strip()
    if not text:
        return Decimal("0")
    return Decimal(text)


def _parse_date(value: str) -> date | None:
    text = (value or "").strip()
    if not text:
        return None
    return date.fromisoformat(text)


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _sheet_setup(ws: Worksheet, header: list[str], widths: list[int] | None = None) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"
    if widths:
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width


def _write_summary(
    ws: Worksheet,
    invoices: list[dict[str, str]],
    eval_rows: list[dict[str, str]],
) -> None:
    decision_counts: dict[str, int] = defaultdict(int)
    grand_total = Decimal("0")
    posted_amount = Decimal("0")
    for row in invoices:
        decision = row.get("decision") or "UNKNOWN"
        decision_counts[decision] += 1
        gt = _parse_decimal(row.get("grand_total", ""))
        grand_total += gt
        if (row.get("posted") or "").lower() == "true":
            posted_amount += gt

    macro_f1 = ""
    for row in eval_rows:
        if row.get("metric") == "header_macro_f1":
            macro_f1 = row.get("value") or ""
            break

    header = ["Metric", "Value"]
    _sheet_setup(ws, header, [36, 18])
    rows = [
        ("Invoice count (ALLOW)", decision_counts.get("ALLOW", 0)),
        ("Invoice count (REVIEW)", decision_counts.get("REVIEW", 0)),
        ("Invoice count (BLOCK)", decision_counts.get("BLOCK", 0)),
        ("Grand total (all invoices)", grand_total),
        ("Posted amount", posted_amount),
        ("Extraction macro F1", macro_f1),
    ]
    for metric, value in rows:
        if isinstance(value, Decimal):
            ws.append([metric, float(value)])
            ws.cell(row=ws.max_row, column=2).number_format = INR_FORMAT
        elif isinstance(value, int):
            ws.append([metric, value])
            ws.cell(row=ws.max_row, column=2).number_format = _INT_FORMAT
        else:
            ws.append([metric, value])


def _aging_bucket(days: int) -> str:
    if days <= 30:
        return "0-30"
    if days <= 60:
        return "31-60"
    if days <= 90:
        return "61-90"
    return ">90"


def _write_aging(ws: Worksheet, invoices: list[dict[str, str]], as_of: date) -> None:
    header = ["Bucket", "Invoice count", "Unposted amount"]
    _sheet_setup(ws, header, [12, 16, 20])
    buckets: dict[str, tuple[int, Decimal]] = {
        "0-30": (0, Decimal("0")),
        "31-60": (0, Decimal("0")),
        "61-90": (0, Decimal("0")),
        ">90": (0, Decimal("0")),
    }
    for row in invoices:
        if (row.get("posted") or "").lower() == "true":
            continue
        inv_date = _parse_date(row.get("invoice_date", ""))
        if inv_date is None:
            continue
        days = (as_of - inv_date).days
        if days < 0:
            continue
        bucket = _aging_bucket(days)
        count, amount = buckets[bucket]
        buckets[bucket] = (count + 1, amount + _parse_decimal(row.get("grand_total", "")))

    for bucket in ("0-30", "31-60", "61-90", ">90"):
        count, amount = buckets[bucket]
        ws.append([bucket, count, float(amount)])
        ws.cell(row=ws.max_row, column=2).number_format = _INT_FORMAT
        ws.cell(row=ws.max_row, column=3).number_format = INR_FORMAT


def _write_gst_by_rate(ws: Worksheet, lines: list[dict[str, str]]) -> None:
    header = ["gst_rate", "taxable_value", "cgst", "sgst", "igst", "cess"]
    _sheet_setup(ws, header, [10, 18, 14, 14, 14, 14])
    totals: dict[str, dict[str, Decimal]] = defaultdict(
        lambda: {
            "taxable_value": Decimal("0"),
            "cgst": Decimal("0"),
            "sgst": Decimal("0"),
            "igst": Decimal("0"),
            "cess": Decimal("0"),
        }
    )
    for row in lines:
        rate = row.get("gst_rate") or "0.00"
        bucket = totals[rate]
        bucket["taxable_value"] += _parse_decimal(row.get("taxable_value", ""))
        bucket["cgst"] += _parse_decimal(row.get("cgst", ""))
        bucket["sgst"] += _parse_decimal(row.get("sgst", ""))
        bucket["igst"] += _parse_decimal(row.get("igst", ""))
        bucket["cess"] += _parse_decimal(row.get("cess", ""))

    for rate in sorted(totals, key=lambda r: _parse_decimal(r)):
        bucket = totals[rate]
        ws.append(
            [
                rate,
                float(bucket["taxable_value"]),
                float(bucket["cgst"]),
                float(bucket["sgst"]),
                float(bucket["igst"]),
                float(bucket["cess"]),
            ]
        )
        for col in range(2, 7):
            ws.cell(row=ws.max_row, column=col).number_format = INR_FORMAT


def _write_vendors(ws: Worksheet, invoices: list[dict[str, str]]) -> None:
    header = ["vendor_name", "vendor_id", "invoice_count", "grand_total", "unposted_amount"]
    _sheet_setup(ws, header, [28, 12, 14, 18, 18])
    stats: dict[str, dict[str, Decimal | int | str]] = defaultdict(
        lambda: {
            "vendor_id": "",
            "invoice_count": 0,
            "grand_total": Decimal("0"),
            "unposted_amount": Decimal("0"),
        }
    )
    for row in invoices:
        name = row.get("vendor_name") or row.get("vendor_gstin") or "(unknown)"
        entry = stats[name]
        if not entry["vendor_id"]:
            entry["vendor_id"] = row.get("vendor_id") or ""
        entry["invoice_count"] = int(entry["invoice_count"]) + 1
        gt = _parse_decimal(row.get("grand_total", ""))
        entry["grand_total"] = Decimal(entry["grand_total"]) + gt
        if (row.get("posted") or "").lower() != "true":
            entry["unposted_amount"] = Decimal(entry["unposted_amount"]) + gt

    for name in sorted(stats):
        entry = stats[name]
        ws.append(
            [
                name,
                entry["vendor_id"],
                int(entry["invoice_count"]),
                float(Decimal(entry["grand_total"])),
                float(Decimal(entry["unposted_amount"])),
            ]
        )
        ws.cell(row=ws.max_row, column=3).number_format = _INT_FORMAT
        for col in (4, 5):
            ws.cell(row=ws.max_row, column=col).number_format = INR_FORMAT


def _write_exceptions(ws: Worksheet, invoices: list[dict[str, str]]) -> None:
    header = ["decision", "top_issue_code", "invoice_count", "exposure"]
    _sheet_setup(ws, header, [12, 28, 14, 18])
    groups: dict[tuple[str, str], tuple[int, Decimal]] = defaultdict(lambda: (0, Decimal("0")))
    for row in invoices:
        decision = row.get("decision") or ""
        if decision not in {"REVIEW", "BLOCK"}:
            continue
        key = (decision, row.get("top_issue_code") or "")
        count, exposure = groups[key]
        groups[key] = (count + 1, exposure + _parse_decimal(row.get("grand_total", "")))

    for (decision, issue), (count, exposure) in sorted(groups.items()):
        ws.append([decision, issue, count, float(exposure)])
        ws.cell(row=ws.max_row, column=3).number_format = _INT_FORMAT
        ws.cell(row=ws.max_row, column=4).number_format = INR_FORMAT


def build_mis_workbook(data_dir: Path, out_path: Path, *, as_of: date | None = None) -> None:
    as_of_date = as_of or date.today()
    invoices = _read_csv(data_dir / "invoices.csv")
    lines = _read_csv(data_dir / "line_items.csv")
    eval_rows = _read_csv(data_dir / "eval_metrics.csv")

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    writers = {
        "Summary": lambda ws: _write_summary(ws, invoices, eval_rows),
        "Aging": lambda ws: _write_aging(ws, invoices, as_of_date),
        "GST by rate": lambda ws: _write_gst_by_rate(ws, lines),
        "Vendors": lambda ws: _write_vendors(ws, invoices),
        "Exceptions": lambda ws: _write_exceptions(ws, invoices),
    }
    for sheet_name in MIS_SHEETS:
        ws = wb.create_sheet(sheet_name)
        writers[sheet_name](ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build Excel MIS from BI CSV marts")
    parser.add_argument("--data", type=Path, required=True, help="Directory with exported CSVs")
    parser.add_argument("--out", type=Path, required=True, help="Output .xlsx path")
    parser.add_argument(
        "--as-of",
        type=str,
        default=None,
        help="Aging reference date (YYYY-MM-DD); default today",
    )
    args = parser.parse_args(argv)
    as_of = date.fromisoformat(args.as_of) if args.as_of else None
    build_mis_workbook(args.data, args.out, as_of=as_of)
    print(f"Wrote MIS workbook to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
