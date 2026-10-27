import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from bi.mis import MIS_SHEETS, build_mis_workbook

DATA_DIR = Path(__file__).resolve().parents[2] / "bi" / "data"
AS_OF = date(2026, 10, 26)


def _read_csv(name: str) -> list[dict[str, str]]:
    path = DATA_DIR / name
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _dec(value: str) -> Decimal:
    return Decimal((value or "0").strip() or "0")


def test_mis_workbook_against_committed_csvs(tmp_path: Path) -> None:
    out = tmp_path / "mis.xlsx"
    build_mis_workbook(DATA_DIR, out, as_of=AS_OF)
    wb = load_workbook(out, read_only=True, data_only=True)

    assert wb.sheetnames == list(MIS_SHEETS)

    invoices = _read_csv("invoices.csv")
    lines = _read_csv("line_items.csv")
    eval_rows = _read_csv("eval_metrics.csv")

    summary = wb["Summary"]
    allow = sum(1 for r in invoices if r.get("decision") == "ALLOW")
    review = sum(1 for r in invoices if r.get("decision") == "REVIEW")
    block = sum(1 for r in invoices if r.get("decision") == "BLOCK")
    posted = sum(_dec(r["grand_total"]) for r in invoices if r.get("posted") == "true")
    macro = next(r["value"] for r in eval_rows if r["metric"] == "header_macro_f1")

    summary_map = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2, max_col=2)}
    assert summary_map["Invoice count (ALLOW)"] == allow
    assert summary_map["Invoice count (REVIEW)"] == review
    assert summary_map["Invoice count (BLOCK)"] == block
    assert Decimal(str(summary_map["Posted amount"])) == posted
    assert str(summary_map["Extraction macro F1"]) == macro

    unposted = [r for r in invoices if r.get("posted") != "true"]
    buckets: dict[str, int] = defaultdict(int)
    for row in unposted:
        raw_date = (row.get("invoice_date") or "").strip()
        if not raw_date:
            continue
        inv_date = date.fromisoformat(raw_date)
        days = (AS_OF - inv_date).days
        if days <= 30:
            buckets["0-30"] += 1
        elif days <= 60:
            buckets["31-60"] += 1
        elif days <= 90:
            buckets["61-90"] += 1
        else:
            buckets[">90"] += 1
    aging = wb["Aging"]
    for row in aging.iter_rows(min_row=2, max_col=2):
        if row[0].value:
            assert row[1].value == buckets[row[0].value]

    rate_totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for row in lines:
        rate_totals[row["gst_rate"]] += _dec(row["cgst"])
    gst = wb["GST by rate"]
    gst_rows = {
        str(r[0].value): Decimal(str(r[2].value or 0))
        for r in gst.iter_rows(min_row=2, max_col=3)
        if r[0].value is not None
    }
    for rate, cgst in rate_totals.items():
        assert gst_rows[rate] == cgst

    wb.close()
